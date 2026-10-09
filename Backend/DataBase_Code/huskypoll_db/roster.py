"""Roster import: the StudentAdmin class-roster spreadsheet, or a pasted list of IDs/emails."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Course, Enrollment, User
from .services import STUDENT_ID_RE, HuskyPollError, enroll, find_or_create_user, log_event

# StudentAdmin column headers -> our field names (matched case-insensitively).
STUDENTADMIN_COLUMNS = {
    "id": "student_id",
    "name": "name",
    "netid": "net_id",
    "email address": "email",
    "section": "section",
}


@dataclass
class ImportResult:
    users_created: int = 0
    users_matched: int = 0
    enrolled: int = 0
    reactivated: int = 0
    unchanged: int = 0
    dropped: int = 0
    skipped: list[tuple[int, str]] = field(default_factory=list)  # (row number, reason)

    def summary(self) -> str:
        s = (f"{self.enrolled} added, {self.reactivated} re-activated, {self.unchanged} already enrolled, "
             f"{self.dropped} dropped; {self.users_created} new users, {self.users_matched} existing")
        if self.skipped:
            s += f"; {len(self.skipped)} rows skipped"
        return s


def _split_name(name: Optional[str]) -> tuple[Optional[str], Optional[str]]:
    """StudentAdmin writes 'LastName,FirstName'."""
    if not name:
        return None, None
    if "," in name:
        last, first = name.split(",", 1)
        return first.strip() or None, last.strip() or None
    parts = name.split()
    return (" ".join(parts[:-1]) or None, parts[-1]) if len(parts) > 1 else (None, parts[0])


def read_studentadmin_xlsx(path: str | Path) -> list[dict]:
    """Read the roster rows as dicts with keys student_id, name, net_id, email, section, _row."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header_map: dict[int, str] = {}
    header_row = 0
    for header_row, row in enumerate(rows, start=1):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if "id" in cells and "name" in cells:
            header_map = {i: STUDENTADMIN_COLUMNS[c] for i, c in enumerate(cells) if c in STUDENTADMIN_COLUMNS}
            break
        if header_row > 20:
            break
    if not header_map:
        raise HuskyPollError("could not find the header row (expected columns 'ID' and 'Name')")

    out = []
    for rownum, row in enumerate(rows, start=header_row + 1):
        rec = {key: row[i] for i, key in header_map.items() if i < len(row)}
        if not any(v not in (None, "") for v in rec.values()):
            continue  # blank line
        rec = {k: (str(v).strip() if v is not None else None) for k, v in rec.items()}
        sid = rec.get("student_id")
        if sid and sid.endswith(".0"):  # Excel stored the ID as a number
            sid = sid[:-2]
        rec["student_id"] = sid.zfill(7) if sid and sid.isdigit() else sid
        rec["_row"] = rownum
        out.append(rec)
    wb.close()
    return out


def import_records(
    db: Session,
    course: Course,
    records: Iterable[dict],
    *,
    drop_missing: bool = False,
    by: Optional[User] = None,
) -> ImportResult:
    """Create/match users and enroll them. With drop_missing, students no longer on
    the roster are marked dropped (their past responses are kept)."""
    result = ImportResult()
    seen: set[int] = set()
    for rec in records:
        first, last = _split_name(rec.get("name"))
        try:
            user, created = find_or_create_user(
                db,
                net_id=rec.get("net_id"),
                email=rec.get("email"),
                student_id=rec.get("student_id"),
                first_name=rec.get("first_name") or first,
                last_name=rec.get("last_name") or last,
            )
        except HuskyPollError as exc:
            result.skipped.append((rec.get("_row", 0), str(exc)))
            continue
        result.users_created += created
        result.users_matched += not created
        _, outcome = enroll(db, course, user, source="import", section=rec.get("section"))
        result.enrolled += outcome == "added"
        result.reactivated += outcome == "reactivated"
        result.unchanged += outcome == "unchanged"
        seen.add(user.id)

    if drop_missing:
        for enr in db.scalars(
            select(Enrollment).where(Enrollment.course_id == course.id, Enrollment.status == "active")
        ):
            if enr.user_id not in seen:
                enr.status = "dropped"
                result.dropped += 1
    db.flush()
    log_event(db, "roster_import", user=by, course=course, summary=result.summary(),
              skipped=result.skipped or None)
    return result


def import_studentadmin_xlsx(
    db: Session, course: Course, path: str | Path, *, drop_missing: bool = False, by: Optional[User] = None
) -> ImportResult:
    return import_records(db, course, read_studentadmin_xlsx(path), drop_missing=drop_missing, by=by)


def import_identifiers(
    db: Session, course: Course, text: str, *, by: Optional[User] = None
) -> ImportResult:
    """Import a pasted list: one NetID, email or 7-digit student ID per line (commas also work)."""
    records = []
    for n, token in enumerate(text.replace(",", "\n").splitlines(), start=1):
        token = token.strip()
        if not token:
            continue
        if "@" in token:
            records.append({"email": token, "_row": n})
        elif STUDENT_ID_RE.match(token):
            records.append({"student_id": token, "_row": n})
        else:
            records.append({"net_id": token, "_row": n})  # an invalid NetID is reported in .skipped
    return import_records(db, course, records, by=by)
