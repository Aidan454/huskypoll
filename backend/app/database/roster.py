"""Roster import: the StudentAdmin class roster (.xlsx or .csv), or a pasted list of IDs/emails.

A StudentAdmin import is the whole class list, so it also reports who is on the
roster here but missing from the file (usually students who dropped the course).
Those students are not dropped automatically: the instructor reviews the list
and drops them with drop_students().
"""
from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Course, Enrollment, User
from .services import (
    STUDENT_ID_RE,
    HuskyPollError,
    enroll,
    find_or_create_user,
    log_event,
)

# StudentAdmin column headers -> our field names (matched case-insensitively).
STUDENTADMIN_COLUMNS = {
    "id": "student_id",
    "student id": "student_id",
    "name": "name",
    "netid": "net_id",
    "net id": "net_id",
    "email": "email",
    "email address": "email",
    "section": "section",
}


@dataclass
class ImportResult:
    users_created: int = 0
    users_matched: int = 0
    enrolled: int = 0
    reactivated: int = 0
    unchanged: int = 0
    dropped: int = 0
    skipped: list[tuple[int, str]] = field(default_factory=list)  # (row number, reason)
    added: list[User] = field(default_factory=list)  # newly enrolled or re-activated
    missing: list[User] = field(default_factory=list)  # active here, not in the file (full imports only)

    def summary(self) -> str:
        s = (f"{self.enrolled} added, {self.reactivated} re-activated, {self.unchanged} already enrolled, "
             f"{self.dropped} dropped; {self.users_created} new users, {self.users_matched} existing")
        if self.skipped:
            s += f"; {len(self.skipped)} rows skipped"
        if self.missing:
            s += f"; {len(self.missing)} on the roster but not in the file"
        return s


def _split_name(name: str | None) -> tuple[str | None, str | None]:
    """StudentAdmin writes 'LastName,FirstName'."""
    if not name:
        return None, None
    if "," in name:
        last, first = name.split(",", 1)
        return first.strip() or None, last.strip() or None
    parts = name.split()
    return (" ".join(parts[:-1]) or None, parts[-1]) if len(parts) > 1 else (None, parts[0])


def _read_rows(rows: Iterable[Iterable]) -> list[dict]:
    """Find the header row, then return the rows below it as dicts with keys
    student_id, name, net_id, email, section, _row.

    NetIDs come back in UPPERCASE and emails in lowercase, however the file wrote them.
    """
    rows = iter(rows)
    header_map: dict[int, str] = {}
    header_row = 0
    for header_row, row in enumerate(rows, start=1):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        found = {STUDENTADMIN_COLUMNS[c] for c in cells if c in STUDENTADMIN_COLUMNS}
        if {"student_id", "name"} <= found or found & {"net_id", "email"}:
            header_map = {i: STUDENTADMIN_COLUMNS[c] for i, c in enumerate(cells) if c in STUDENTADMIN_COLUMNS}
            break
        if header_row > 20:
            break
    if not header_map:
        raise HuskyPollError("could not find the header row (expected columns 'NetID' and 'Email', or 'ID' and 'Name')")

    out = []
    for rownum, row in enumerate(rows, start=header_row + 1):
        row = list(row)
        rec = {key: row[i] for i, key in header_map.items() if i < len(row)}
        if not any(v not in (None, "") for v in rec.values()):
            continue  # blank line
        rec = {k: (str(v).strip() or None if v is not None else None) for k, v in rec.items()}
        sid = rec.get("student_id")
        if sid and sid.endswith(".0"):  # Excel stored the ID as a number
            sid = sid[:-2]
        rec["student_id"] = sid.zfill(7) if sid and sid.isdigit() else sid
        if rec.get("net_id"):
            rec["net_id"] = rec["net_id"].upper()
        if rec.get("email"):
            rec["email"] = rec["email"].lower()
        rec["_row"] = rownum
        out.append(rec)
    return out


def read_studentadmin_xlsx(path: str | Path) -> list[dict]:
    """Read the roster rows from the first sheet of a StudentAdmin .xlsx."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return _read_rows(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()


def read_studentadmin_csv(data: bytes | str) -> list[dict]:
    """Read the roster rows from a StudentAdmin .csv (e.g. columns NetID, Email)."""
    if isinstance(data, bytes):
        try:
            data = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            data = data.decode("cp1252")  # Excel on Windows
    return _read_rows(csv.reader(io.StringIO(data.lstrip("\ufeff"))))


def import_records(
    db: Session,
    course: Course,
    records: Iterable[dict],
    *,
    drop_missing: bool = False,
    report_missing: bool = False,
    by: User | None = None,
) -> ImportResult:
    """Create/match users and enroll them.

    With report_missing, `missing` lists active students who are not in `records`.
    With drop_missing, they are also marked dropped (their past responses are kept).
    """
    result = ImportResult()
    seen: set[int] = set()
    for rec in records:
        first, last = _split_name(rec.get("name"))
        try:
            user, created = find_or_create_user(
                db,
                net_id=rec.get("net_id"),
                email=rec.get("email"),
                student_id=rec.get("student_id"),
                first_name=rec.get("first_name") or first,
                last_name=rec.get("last_name") or last,
            )
        except HuskyPollError as exc:
            result.skipped.append((rec.get("_row", 0), str(exc)))
            continue
        result.users_created += created
        result.users_matched += not created
        _, outcome = enroll(db, course, user, source="import", section=rec.get("section"))
        result.enrolled += outcome == "added"
        result.reactivated += outcome == "reactivated"
        result.unchanged += outcome == "unchanged"
        if outcome != "unchanged" and user.id not in seen:
            result.added.append(user)
        seen.add(user.id)

    if drop_missing or report_missing:
        for enr in db.scalars(
            select(Enrollment).where(Enrollment.course_id == course.id, Enrollment.status == "active")
        ):
            if enr.user_id not in seen:
                result.missing.append(enr.user)
                if drop_missing:
                    enr.status = "dropped"
                    result.dropped += 1
        result.missing.sort(key=lambda u: ((u.last_name or "").lower(), (u.first_name or "").lower(), u.net_id or ""))
    db.flush()
    log_event(db, "roster_import", user=by, course=course, summary=result.summary(),
              skipped=result.skipped or None)
    return result


def import_studentadmin_xlsx(
    db: Session, course: Course, path: str | Path, *, drop_missing: bool = False, by: User | None = None
) -> ImportResult:
    return import_records(db, course, read_studentadmin_xlsx(path),
                          drop_missing=drop_missing, report_missing=True, by=by)


def import_studentadmin_csv(
    db: Session, course: Course, data: bytes | str, *, drop_missing: bool = False, by: User | None = None
) -> ImportResult:
    return import_records(db, course, read_studentadmin_csv(data),
                          drop_missing=drop_missing, report_missing=True, by=by)


def drop_students(db: Session, course: Course, user_ids: Iterable[int], *, by: User | None = None) -> int:
    """Drop the given students (e.g. the `missing` list from an import). Returns how many were dropped.

    Ids that are not active students in the course are ignored. Past responses are kept.
    """
    ids = set(user_ids)
    if not ids:
        return 0
    enrollments = db.scalars(select(Enrollment).where(
        Enrollment.course_id == course.id, Enrollment.status == "active", Enrollment.user_id.in_(ids)
    )).all()
    for enr in enrollments:
        enr.status = "dropped"
    db.flush()
    log_event(db, "roster_drop_missing", user=by, course=course,
              dropped_user_ids=sorted(e.user_id for e in enrollments))
    return len(enrollments)


def import_identifiers(
    db: Session, course: Course, text: str, *, by: User | None = None
) -> ImportResult:
    """Import a pasted list: one NetID, email or 7-digit student ID per line (commas also work)."""
    records = []
    for n, token in enumerate(text.replace(",", "\n").splitlines(), start=1):
        token = token.strip()
        if not token:
            continue
        if "@" in token:
            records.append({"email": token.lower(), "_row": n})
        elif STUDENT_ID_RE.match(token):
            records.append({"student_id": token, "_row": n})
        else:
            records.append({"net_id": token.upper(), "_row": n})  # an invalid NetID is reported in .skipped
    return import_records(db, course, records, by=by)
